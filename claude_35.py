"""

Sample bot using Claude-3.5-Haiku with file attachment support.

"""

from __future__ import annotations

from typing import AsyncIterable
from docx import Document  # Для работы с .docx
from PyPDF2 import PdfReader  # Для работы с PDF
import openpyxl  # Для работы с .xlsx

import fastapi_poe as fp
from modal import App, Image, asgi_app


class ClaudeBot(fp.PoeBot):
    async def get_response(
        self, request: fp.QueryRequest
    ) -> AsyncIterable[fp.PartialResponse]:
        responses = []

        # Проверяем вложения внутри query
        for message in request.query:
            if message.role == "attachment" and message.metadata:
                attachment_name = message.metadata.get("name", "Без имени")
                if attachment_name.endswith(".docx"):
                    try:
                        # Обработка .docx файла
                        doc = Document(message.content)
                        doc_text = "\n".join([p.text for p in doc.paragraphs])
                        responses.append(f"Текст из .docx файла {attachment_name}:\n{doc_text}")
                    except Exception as e:
                        responses.append(f"Ошибка обработки .docx файла {attachment_name}: {str(e)}")
                elif attachment_name.endswith(".pdf"):
                    try:
                        # Обработка PDF файла
                        pdf_reader = PdfReader(message.content)
                        pdf_text = "\n".join(page.extract_text() for page in pdf_reader.pages)
                        responses.append(f"Текст из PDF файла {attachment_name}:\n{pdf_text}")
                    except Exception as e:
                        responses.append(f"Ошибка обработки PDF файла {attachment_name}: {str(e)}")
                elif attachment_name.endswith(".xlsx"):
                    try:
                        # Обработка .xlsx файла
                        workbook = openpyxl.load_workbook(message.content, data_only=True)
                        sheet = workbook.active
                        rows = []
                        for row in sheet.iter_rows(values_only=True):
                            rows.append("\t".join(str(cell) if cell is not None else "" for cell in row))
                        xlsx_text = "\n".join(rows)
                        responses.append(f"Содержимое .xlsx файла {attachment_name}:\n{xlsx_text}")
                    except Exception as e:
                        responses.append(f"Ошибка обработки .xlsx файла {attachment_name}: {str(e)}")

        # Если нет вложений, отправляем запрос только Claude-3.5-Haiku
        if not responses:
            async for msg in fp.stream_request(request, "Claude-3.5-Haiku", request.access_key):
                yield msg
        else:
            # Возвращаем обработанные вложения
            yield fp.PartialResponse(text="\n".join(responses))

    async def get_settings(self, setting: fp.SettingsRequest) -> fp.SettingsResponse:
        return fp.SettingsResponse(
            allow_attachments=True,
            expand_text_attachments=True,
            enable_image_comprehension=True,
            server_bot_dependencies={"Claude-3.5-Haiku": 1},
        )


REQUIREMENTS = ["fastapi-poe==0.0.48", "python-docx", "PyPDF2", "openpyxl"]
image = Image.debian_slim().pip_install(*REQUIREMENTS)
app = App("claude-bot")


@app.function(image=image)
@asgi_app()
def fastapi_app():
    bot = ClaudeBot()
    app = fp.make_app(bot, allow_without_key=True)
    return app
